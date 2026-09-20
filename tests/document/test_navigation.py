"""Document order, typed path resolution, and anchored comment lookup."""

from docx import Document
import pytest

from docxnote import Cell, Comment, DocxDocument, Paragraph, Table
from tests.support.docx import save_docx


@pytest.fixture
def document():
    source = Document()
    source.add_paragraph("first")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    source.add_paragraph("")
    source.add_table(rows=2, cols=1)
    source.add_paragraph("last")
    return DocxDocument.parse(save_docx(source))


def test_blocks_preserve_order_empty_paragraphs_and_independent_indices(document):
    blocks = document.blocks()
    assert isinstance(blocks, tuple)
    assert [(type(block), block.path) for block in blocks] == [
        (Paragraph, "p:0"),
        (Table, "t:0"),
        (Paragraph, "p:1"),
        (Table, "t:1"),
        (Paragraph, "p:2"),
    ]
    assert [p.text for p in blocks if isinstance(p, Paragraph)] == ["first", "", "last"]
    assert [t.shape() for t in blocks if isinstance(t, Table)] == [(1, 2), (2, 1)]
    assert [(p.path, p.text) for p in document.iter_paragraphs()] == [
        ("p:0", "first"),
        ("t:0/r:0/c:0/p:0", "left"),
        ("t:0/r:0/c:1/p:0", "right"),
        ("p:1", ""),
        ("t:1/r:0/c:0/p:0", ""),
        ("t:1/r:1/c:0/p:0", ""),
        ("p:2", "last"),
    ]


@pytest.mark.parametrize(
    "path,kind",
    [
        ("p:2", Paragraph),
        ("t:0", Table),
        ("t:0/r:0/c:1", Cell),
        ("t:0/r:0/c:1/p:0", Paragraph),
    ],
)
def test_resolve_returns_the_addressed_view(document, path, kind):
    view = document.resolve(path)
    assert isinstance(view, kind)
    assert view.path == path


@pytest.mark.parametrize(
    "path,error",
    [
        ("", ValueError),
        ("p:999", LookupError),
        ("t:999", LookupError),
        ("t:0/p:0", ValueError),
        ("t:0/r:0", ValueError),
        ("t:0/r:0/p:0", ValueError),
        ("t:0/r:9/c:0", LookupError),
        ("t:0/r:0/c:0/p:999", LookupError),
        ("t:0/r:0/c:0/t:999", LookupError),
        ("p:0#999", LookupError),
        ("t:0#0", ValueError),
        ("p:0/p:0", ValueError),
        ("r:0/c:0", ValueError),
    ],
)
def test_resolve_distinguishes_invalid_structure_from_missing_objects(
    document, path, error
):
    with pytest.raises(error):
        document.resolve(path)


def test_comments_are_collected_in_document_order_and_paths_roundtrip(document):
    # Write in reverse document order; reads must still follow document order.
    for path in ("p:2", "t:0/r:0/c:1/p:0", "p:0"):
        paragraph = document.resolve(path)
        assert isinstance(paragraph, Paragraph)
        paragraph.comment(path, 0, 2)
    expected_paths = ["p:0#2", "t:0/r:0/c:1/p:0#1", "p:2#0"]
    for current in (
        document,
        DocxDocument.parse(document.render(), keep_comments=True),
    ):
        assert [c.path for c in current.comments()] == expected_paths
        for path in expected_paths:
            comment = current.resolve(path)
            assert isinstance(comment, Comment)
            assert (comment.path, comment.start, comment.end) == (path, 0, 2)
            assert comment.text == comment.paragraph.path
        normalized = current.resolve("  p:0 / #2  ")
        assert isinstance(normalized, Comment) and normalized.path == "p:0#2"
