"""Block content controls are transparent to paths and document traversal."""

import json

from docx import Document
from lxml import etree
import pytest

from docxnote import DocxDocument, DocxShell, Paragraph
from docxnote.namespaces import NS
from tests.support.docx import save_docx


def _wrap(element):
    parent = element.getparent()
    index = parent.index(element)
    sdt = etree.Element(f"{{{NS['w']}}}sdt")
    etree.SubElement(sdt, f"{{{NS['w']}}}sdtPr")
    content = etree.SubElement(sdt, f"{{{NS['w']}}}sdtContent")
    content.append(element)
    parent.insert(index, sdt)
    return sdt


@pytest.mark.parametrize("in_cell", [False, True])
def test_nested_content_controls_expose_paragraphs_and_tables(in_cell):
    source = Document()
    container = source.add_table(rows=1, cols=1).cell(0, 0) if in_cell else source
    if in_cell:
        empty = container.paragraphs[0]._p
        parent = empty.getparent()
        assert parent is not None
        parent.remove(empty)
    _wrap(_wrap(container.add_paragraph("inside")._p))
    table = container.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "nested table"
    _wrap(table._tbl)
    # Cells add a mandatory empty paragraph after a nested table.
    if in_cell:
        empty = container.paragraphs[-1]._p
        parent = empty.getparent()
        assert parent is not None
        parent.remove(empty)
    container.add_paragraph("outside")
    doc = DocxDocument.parse(save_docx(source))
    prefix = "t:0/r:0/c:0/" if in_cell else ""
    expected = [
        (prefix + "p:0", "inside"),
        (prefix + "t:0/r:0/c:0/p:0", "nested table"),
        (prefix + "p:1", "outside"),
    ]
    assert [(p.path, p.text) for p in doc.iter_paragraphs()] == expected
    for path, text in expected:
        paragraph = doc.resolve(path)
        assert isinstance(paragraph, Paragraph)
        assert paragraph.text == text
    paragraph = doc.resolve(prefix + "p:0")
    assert isinstance(paragraph, Paragraph)
    paragraph.comment("control note", 1, 4)
    shell = DocxShell(doc)
    result = shell.run("docx")
    assert result["exit_code"] == 0
    assert [
        (r["path"], r["text"]) for r in map(json.loads, result["stdout"].splitlines())
    ] == expected
    reopened = DocxDocument.parse(doc.render(), keep_comments=True)
    assert [(c.path, c.start, c.end) for c in reopened.comments()] == [
        (prefix + "p:0#0", 1, 4)
    ]
