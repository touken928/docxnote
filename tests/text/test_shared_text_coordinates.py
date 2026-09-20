"""All views keep host text coordinates through repeated run splitting."""

from docx import Document
from lxml import etree

from docxnote import DocxDocument, Paragraph
from docxnote.namespaces import NS
from tests.support.docx import save_docx, xml_part


W = "{" + NS["w"] + "}"


def test_repeated_splits_share_coordinates_across_wrappers_and_preserve_run_content():
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run("A")
    control = etree.Element(W + "sdt")
    paragraph._p.append(control)
    content = etree.SubElement(control, W + "sdtContent")
    hyperlink = etree.SubElement(content, W + "hyperlink", {W + "anchor": "bookmark"})
    run = etree.SubElement(hyperlink, W + "r")
    properties = etree.SubElement(run, W + "rPr")
    etree.SubElement(properties, W + "b")
    etree.SubElement(run, W + "t").text = "B "
    etree.SubElement(run, W + "tab")
    etree.SubElement(run, W + "t").text = "中🙂"
    etree.SubElement(run, W + "br")
    drawing = etree.SubElement(run, W + "drawing")
    box = etree.SubElement(drawing, W + "txbxContent")
    inner_paragraph = etree.SubElement(box, W + "p")
    inner_run = etree.SubElement(inner_paragraph, W + "r")
    etree.SubElement(inner_run, W + "t").text = "EXCLUDED"
    etree.SubElement(run, W + "t").text = "CD"
    paragraph.add_run("Z")
    document = DocxDocument.parse(save_docx(source))
    first = document.resolve("p:0")
    second = document.resolve("p:0")
    assert isinstance(first, Paragraph) and isinstance(second, Paragraph)
    expected_text = "AB \t中🙂\nCDZ"
    assert first.text == second.text == expected_text

    expected_ranges = {}
    for index, (start, end) in enumerate(((2, 8), (4, 6), (7, 8), (3, 3), (0, 10))):
        view = first if index % 2 else second
        added = view.comment(f"range {index}", start, end)
        expected_ranges[added.path] = (start, end)
        for wrapper in (first, second):
            assert wrapper.text == expected_text
            assert {
                comment.path: (comment.start, comment.end)
                for comment in wrapper.comments
            } == expected_ranges

    rendered = document.render()
    reopened = DocxDocument.parse(rendered, keep_comments=True)
    assert next(reopened.iter_paragraphs()).text == expected_text
    assert {
        comment.path: (comment.start, comment.end) for comment in reopened.comments()
    } == expected_ranges
    root = xml_part(rendered)
    assert len(root.findall(".//w:drawing", NS)) == 1
    box_text = root.find(".//w:txbxContent/w:p/w:r/w:t", NS)
    assert box_text is not None and box_text.text == "EXCLUDED"
    for run in root.findall(".//w:hyperlink/w:r", NS):
        if run.find("./w:t", NS) is not None:
            assert run.find("./w:rPr/w:b", NS) is not None
