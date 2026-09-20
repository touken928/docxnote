"""Paragraph text ignores formatting and retains every visible character."""

from io import BytesIO

from docx import Document
import pytest

from docxnote import DocxDocument
from tests.support.docx import save_docx


@pytest.mark.parametrize(
    "runs",
    [
        pytest.param(["普通 ", "粗体", " ", "斜体"], id="formatting"),
        pytest.param(["第一行", "\n", "第二行"], id="break"),
        pytest.param(["列1", "\t", "列2"], id="tab"),
        pytest.param(["<>&\"' 中文 🎉"], id="unicode-and-escaping"),
        pytest.param([""], id="empty"),
        pytest.param(["  前导", "    多个空格", "尾随  "], id="whitespace"),
    ],
)
def test_text_matches_python_docx_before_and_after_render(runs):
    source = Document()
    paragraph = source.add_paragraph()
    for text in runs:
        run = paragraph.add_run(text)
        run.bold = True
        run.italic = True
    data = save_docx(source)
    expected = [p.text for p in Document(BytesIO(data)).paragraphs]
    assert expected == ["".join(runs)]
    document = DocxDocument.parse(data)
    assert [p.text for p in document.iter_paragraphs()] == expected
    assert [
        p.text for p in DocxDocument.parse(document.render()).iter_paragraphs()
    ] == expected
