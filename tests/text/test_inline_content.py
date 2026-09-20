"""Inline containers and non-text run children survive partial annotations."""

from docx import Document
from lxml import etree
import pytest

from docxnote import DocxDocument
from docxnote.namespaces import NS
from tests.support.docx import save_docx, xml_part

W = "{" + NS["w"] + "}"


@pytest.mark.parametrize(
    "prefix,text,start,end,expected",
    [
        ("", "ClickHere", 0, 5, (0, 5)),
        ("AB", "CDEF", -4, 999, (2, 6)),
    ],
)
def test_hyperlink_partial_range(prefix, text, start, end, expected):
    source = Document()
    paragraph = source.add_paragraph(prefix)
    relationship = paragraph.part.relate_to(
        "https://example.com",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = etree.Element(W + "hyperlink", {f"{{{NS['r']}}}id": relationship})
    run = etree.SubElement(hyperlink, W + "r")
    etree.SubElement(run, W + "t").text = text
    paragraph._p.append(hyperlink)
    document = DocxDocument.parse(save_docx(source))
    paragraph_view = next(document.iter_paragraphs())
    assert paragraph_view.text == prefix + text
    created = paragraph_view.comment("link", start, end)
    assert (created.start, created.end) == expected
    output = document.render()
    (saved,) = DocxDocument.parse(output, keep_comments=True).comments()
    assert (saved.start, saved.end) == expected
    assert saved.paragraph.text == prefix + text
    assert xml_part(output).find(".//w:hyperlink", NS) is not None


def test_partial_comment_preserves_symbol_attributes():
    source = Document()
    run = source.add_paragraph().add_run("ABCD")
    symbol = etree.Element(W + "sym", {W + "font": "Wingdings", W + "char": "F04A"})
    run._r.append(symbol)
    document = DocxDocument.parse(save_docx(source))
    next(document.iter_paragraphs()).comment("inner", 1, 3)
    output = document.render()
    symbols = xml_part(output).findall(".//w:sym", NS)
    assert len(symbols) == 1
    assert dict(symbols[0].attrib) == dict(symbol.attrib)
    (saved,) = DocxDocument.parse(output, keep_comments=True).comments()
    assert (saved.start, saved.end, saved.paragraph.text) == (1, 3, "ABCD")
