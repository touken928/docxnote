"""Embedded text-box paragraphs must not change their host's coordinates."""

from docx import Document
from lxml import etree
import pytest

from docxnote import DocxDocument, Paragraph, Comment
from docxnote.namespaces import NS
from tests.support.docx import save_docx, xml_part


@pytest.mark.parametrize("kind", ["pict", "drawing"])
def test_textbox_is_preserved_but_excluded_from_host_text_and_ranges(kind):
    source = Document()
    paragraph = source.add_paragraph()
    run = paragraph.add_run("AB")
    drawing = etree.SubElement(run._r, f"{{{NS['w']}}}{kind}")
    if kind == "pict":
        shape = etree.SubElement(drawing, "{urn:schemas-microsoft-com:vml}shape")
        box = etree.SubElement(shape, "{urn:schemas-microsoft-com:vml}textbox")
    else:
        box = etree.SubElement(
            drawing,
            "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}txbx",
        )
    content = etree.SubElement(box, f"{{{NS['w']}}}txbxContent")
    inner = etree.SubElement(content, f"{{{NS['w']}}}p")
    inner_run = etree.SubElement(inner, f"{{{NS['w']}}}r")
    etree.SubElement(inner_run, f"{{{NS['w']}}}t").text = "BOX"
    etree.SubElement(run._r, f"{{{NS['w']}}}t").text = "YZ"
    doc = DocxDocument.parse(save_docx(source))
    paragraph = doc.resolve("p:0")
    assert isinstance(paragraph, Paragraph)
    assert paragraph.text == "ABYZ"
    comment = paragraph.comment("after box", 2, 3)
    resolved = doc.resolve(comment.path)
    assert isinstance(resolved, Comment)
    assert (resolved.start, resolved.end) == (2, 3)
    rendered = doc.render()
    reopened = DocxDocument.parse(rendered, keep_comments=True)
    reopened_paragraph = reopened.resolve("p:0")
    assert isinstance(reopened_paragraph, Paragraph)
    assert reopened_paragraph.text == "ABYZ"
    assert [(c.start, c.end) for c in reopened.comments()] == [(2, 3)]
    root = xml_part(rendered)
    boxes = root.findall(".//w:txbxContent", NS)
    assert len(boxes) == 1
    text = boxes[0].find(".//w:t", NS)
    assert text is not None and text.text == "BOX"
    assert not boxes[0].findall(".//w:commentRangeStart", NS)
