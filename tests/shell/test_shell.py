"""The simulated shell operates on document records, never the host shell."""

import json
import shlex
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from docx import Document

from docxnote import DOCX_SHELL_INSTRUCTIONS, DocxDocument, DocxShell


def test_run_docstring_is_the_exported_tool_contract():
    assert DocxShell.run.__doc__ == DOCX_SHELL_INSTRUCTIONS
    assert "docx [PATH]" in DocxShell.run.__doc__


def make_doc(*texts):
    source = Document()
    for text in texts:
        source.add_paragraph(text)
    stream = BytesIO()
    source.save(stream)
    return DocxDocument.parse(stream.getvalue())


def records(result):
    assert result["exit_code"] == 0, result
    return [json.loads(line) for line in result["stdout"].splitlines()]


def test_sources_and_literal_pipelines():
    shell = DocxShell(make_doc("付款\n期限", "", "Payment [due]", "付款 示例"))
    assert [r["text"] for r in records(shell.run("docx"))] == [
        "付款\n期限",
        "",
        "Payment [due]",
        "付款 示例",
    ]
    assert [
        r["path"] for r in records(shell.run("docx | grep -F '付款' | grep -v '示例'"))
    ] == ["p:0"]
    assert records(shell.run("docx | grep -i payment"))[0]["path"] == "p:2"
    assert records(shell.run("docx | grep '[due]'"))[0]["path"] == "p:2"
    assert len(records(shell.run("docx | grep -e 付款 -e Payment"))) == 3
    assert shell.run("docx | grep -n Payment")["stdout"].startswith("3:")
    assert records(shell.run("docx | head -n 3 | tail -n 2"))[0]["path"] == "p:1"
    assert records(shell.run("docx | head -n 0")) == []
    assert records(shell.run("docx | tail -n 0")) == []
    assert shell.run("docx | grep missing")["exit_code"] == 0


def test_scopes_and_slices():
    source = Document()
    source.add_paragraph("outside")
    for index in range(11):
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f"table {index}"
        table.cell(0, 1).text = "付款🙂\n期限"
    nested = source.tables[1].cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "nested"
    stream = BytesIO()
    source.save(stream)
    shell = DocxShell(DocxDocument.parse(stream.getvalue()))
    rows = records(shell.run("docx t:1"))
    assert all(r["path"].startswith("t:1/") for r in rows)
    assert any(r["text"] == "nested" for r in rows)
    assert records(shell.run("docx t:1/r:0/c:1"))[0]["text"] == "付款🙂\n期限"
    row = records(shell.run("docx t:1/r:0/c:1/p:0 --start 2 --end 4"))[0]
    assert (row["text"], row["start"], row["end"], row["total_chars"]) == (
        "🙂\n",
        2,
        4,
        6,
    )
    assert shell.run("docx t:1 --start 0")["exit_code"] == 2


@pytest.mark.parametrize(
    "command",
    [
        "",
        "wc -l",
        "sed -n 1p",
        "cat file",
        "echo hi",
        "grep hi",
        "docx |",
        "| docx",
        "docx || head",
        "docx ; docx",
        "docx > file",
        "docx && docx",
        "docx\ndocx",
        "docx | grep --unknown hi",
        "docx | head -n -1",
        "docx | tail -n x",
        "docx | head file",
        "docx | grep",
        "docx | grep 'broken",
        "docx p:999",
        "docx p:0#0",
        "docx | comment p:0 hi",
        "comment p:0 hi | head",
        "comment p:0 hi --start 0",
        'comment p:0 hi --quote ""',
        "docx p:0 --start -1",
        "docx p:0 --start 3 --end 2",
    ],
)
def test_errors_do_not_write(command):
    doc = make_doc("hello world")
    result = DocxShell(doc).run(command)
    assert result["exit_code"] == 2
    assert result["stdout"] == ""
    assert result["stderr"]
    assert not doc.comments()


def test_quotes_escaping_and_no_expansion():
    shell = DocxShell(make_doc('a | b; $HOME `id` $(id) say "hi"'))
    for text in ["a | b", "b;", "$HOME", "`id`", "$(id)", 'say "hi"']:
        pattern = shlex.quote(json.dumps(text, ensure_ascii=False)[1:-1])
        assert len(records(shell.run(f"docx | grep -F {pattern}"))) == 1
    result = shell.run('comment p:0 "say \\"hi\\" | keep ; $HOME"')
    assert result["exit_code"] == 0


def test_comment_quote_validation_and_roundtrip():
    doc = make_doc("付款🙂付款")
    shell = DocxShell(doc, author="reviewer")
    for suffix in ["--quote missing", "--quote 付款", "--quote 付款 --start 1"]:
        assert shell.run("comment p:0 note " + suffix)["exit_code"] == 2
    assert not doc.comments()
    added = records(shell.run("comment p:0 '请说明期限' --quote 付款 --start 3"))[0]
    assert (added["target"], added["start"], added["end"], added["author"]) == (
        "p:0",
        3,
        5,
        "reviewer",
    )
    assert records(shell.run("comments p:0"))[0]["path"] == added["path"]
    assert len(shell.added_comments) == 1
    whole = records(shell.run("comment p:0 whole"))[0]
    assert (whole["start"], whole["end"]) == (0, 5)
    reloaded = DocxDocument.parse(doc.render(), keep_comments=True)
    assert len(reloaded.comments()) == 2
    preserved = DocxShell(reloaded)
    assert len(records(preserved.run("comments"))) == 2
    assert preserved.added_comments == ()


def test_overlapping_quote_matches_are_ambiguous():
    assert (
        DocxShell(make_doc("aaa")).run("comment p:0 note --quote aa")["exit_code"] == 2
    )


def test_limit_is_after_filtering_and_json_remains_valid():
    shell = DocxShell(make_doc(*(["x" * 1000] * 40), "needle"), max_output=256)
    result = shell.run("docx | grep needle")
    assert records(result)[0]["path"] == "p:40"
    assert not result["truncated"]
    result = shell.run("docx")
    preview = records(result)[0]
    assert len(result["stdout"]) <= 256
    assert result["truncated"] and preview["text_truncated"]
    assert preview["path"] == "p:0"
    assert preview["text"] == "x" * preview["end"]
    assert preview["total_chars"] == 1000
    next_part = records(
        shell.run(f"docx p:0 --start {preview['end']} --end {preview['end'] + 20}")
    )[0]
    assert next_part["start"] == preview["end"]


def test_record_limit_and_short_circuit(monkeypatch):
    doc = make_doc("a", "b", "c")
    original = doc.iter_paragraphs
    visited = []

    def walk():
        for p in original():
            visited.append(p.path)
            yield p

    monkeypatch.setattr(doc, "iter_paragraphs", walk)
    result = DocxShell(doc).run("docx | head -n 1")
    assert len(records(result)) == 1
    assert visited == ["p:0"]
    assert not result["truncated"]
    result = DocxShell(doc, max_output=256).run("docx")
    assert result["records"] == 3


def test_concurrent_writes_and_snapshot():
    doc = make_doc("hello")
    shell = DocxShell(doc)
    snapshot = shell.added_comments
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(shell.run, [f"comment p:0 note{i}" for i in range(20)]))
    assert all(r["exit_code"] == 0 for r in results)
    assert len(shell.added_comments) == len(doc.comments()) == 20
    assert snapshot == ()


def test_unexpected_errors_propagate(monkeypatch):
    doc = make_doc("hello")

    def broken():
        raise RuntimeError("internal failure")

    monkeypatch.setattr(doc, "iter_paragraphs", broken)
    with pytest.raises(RuntimeError, match="internal failure"):
        DocxShell(doc).run("docx")
