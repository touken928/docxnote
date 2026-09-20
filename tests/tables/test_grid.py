"""Logical dimensions, cell coordinates, and merged-origin identity."""

from docx import Document
import pytest

from docxnote import DocxDocument, Paragraph, Table
from tests.support.docx import save_docx


@pytest.mark.parametrize("rows,cols", [(1, 1), (2, 2), (3, 4), (5, 3), (1, 5), (10, 1)])
def test_rectangular_grid_has_exact_shape_bounds_paths_and_content(rows, cols):
    source = Document()
    original = source.add_table(rows=rows, cols=cols)
    for row in range(rows):
        for col in range(cols):
            original.cell(row, col).text = f"R{row}C{col}"
    document = DocxDocument.parse(save_docx(source))
    table = document.resolve("t:0")
    assert isinstance(table, Table)
    assert table.shape() == (rows, cols)
    for row in range(rows):
        for col in range(cols):
            cell = table[row, col]
            assert cell.bounds() == (row, col, row + 1, col + 1)
            assert cell.path == f"t:0/r:{row}/c:{col}"
            blocks = cell.blocks()
            assert isinstance(blocks, tuple) and len(blocks) == 1
            assert isinstance(blocks[0], Paragraph)
            assert blocks[0].text == original.cell(row, col).text
    for index in [(-1, 0), (0, -1), (rows, 0), (0, cols)]:
        with pytest.raises(IndexError, match="out of bounds"):
            _ = table[index]


def test_mixed_merges_share_origin_bounds_content_and_traversal():
    source = Document()
    original = source.add_table(rows=5, cols=5)
    for row in range(5):
        for col in range(5):
            original.cell(row, col).text = f"R{row}C{col}"
    regions = [(0, 0, 1, 3), (1, 4, 4, 5), (2, 1, 4, 3)]
    for index, (top, left, bottom, right) in enumerate(regions):
        original.cell(top, left).merge(
            original.cell(bottom - 1, right - 1)
        ).text = f"MERGED{index}"
    document = DocxDocument.parse(save_docx(source))
    table = document.resolve("t:0")
    assert isinstance(table, Table)
    assert table.shape() == (5, 5)
    for row in range(5):
        for col in range(5):
            bounds = next(
                (
                    region
                    for region in regions
                    if region[0] <= row < region[2] and region[1] <= col < region[3]
                ),
                (row, col, row + 1, col + 1),
            )
            origin = table[bounds[0], bounds[1]]
            cell = table[row, col]
            assert cell is origin
            assert cell.bounds() == bounds
            assert cell.path == f"t:0/r:{bounds[0]}/c:{bounds[1]}"
            assert [p.text for p in cell.blocks() if isinstance(p, Paragraph)] == [
                original.cell(row, col).text
            ]
            assert document.resolve(f"t:0/r:{row}/c:{col}").path == cell.path
    paths = [p.path for p in document.iter_paragraphs()]
    assert len(paths) == len(set(paths)) == 18
