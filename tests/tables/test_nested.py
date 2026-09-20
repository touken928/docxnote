"""Nested tables preserve block order, direct cell text, paths, and annotations."""

from docx import Document

from docxnote import DocxDocument, Paragraph, Table
from tests.support.docx import save_docx


def test_nested_tables_keep_local_shapes_block_order_and_exact_comment_paths():
    source = Document()
    outer = source.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "plain outer"
    cell = outer.cell(0, 1)
    cell.text = "before"
    inner = cell.add_table(rows=2, cols=2)
    for row in range(2):
        for col in range(2):
            inner.cell(row, col).text = f"inner {row},{col}"
    cell.paragraphs[-1].text = "after"
    deep = inner.cell(1, 0).add_table(rows=1, cols=1)
    deep.cell(0, 0).text = "deep"
    document = DocxDocument.parse(save_docx(source))
    table = document.resolve("t:0")
    assert isinstance(table, Table) and table.shape() == (2, 2)
    assert [p.text for p in table[0, 0].blocks() if isinstance(p, Paragraph)] == [
        "plain outer"
    ]
    blocks = table[0, 1].blocks()
    assert [type(block) for block in blocks] == [Paragraph, Table, Paragraph]
    assert [block.text for block in blocks if isinstance(block, Paragraph)] == [
        "before",
        "after",
    ]
    nested = blocks[1]
    assert isinstance(nested, Table) and nested.shape() == (2, 2)
    assert [p.text for p in nested[0, 1].blocks() if isinstance(p, Paragraph)] == [
        "inner 0,1"
    ]
    path = "t:0/r:0/c:1/t:0/r:1/c:0/t:0/r:0/c:0/p:0"
    paragraph = document.resolve(path)
    assert isinstance(paragraph, Paragraph) and paragraph.text == "deep"
    paragraph.comment("nested note", 1, 3)
    reopened = DocxDocument.parse(document.render(), keep_comments=True)
    (saved,) = reopened.comments()
    assert (saved.path, saved.start, saved.end) == (path + "#0", 1, 3)
    assert [(p.path, p.text) for p in reopened.iter_paragraphs()] == [
        (p.path, p.text) for p in document.iter_paragraphs()
    ]
